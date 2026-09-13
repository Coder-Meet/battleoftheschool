"""Render the public entry and team briefing as offline HTML and editable Word."""

import argparse
import hashlib
from html import escape
from pathlib import Path
import re
import shutil
import subprocess
from urllib.parse import urljoin, urlsplit

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from markdown_it.token import Token

ROOT = Path(__file__).resolve().parents[1]
REPO = "https://github.com/Coder-Meet/battleoftheschool/blob/main/"
DOCUMENTS = ("README", "DEVPOST_SUBMISSION", "PRESENTER_BRIEFING", "DEMO_GUIDE")
TITLES = {
    "DEVPOST_SUBMISSION": "Devpost submission",
    "PRESENTER_BRIEFING": "Presenter briefing",
    "SPEAKER_SCRIPT": "Five-minute presentation script",
}
STYLE = """
body { margin:0; background:#f4f0e7; color:#172b33; font:17px/1.65 Arial,sans-serif; }
main { max-width:1120px; margin:0 auto; padding:55px 60px 90px; }
header { color:#16715f; font-size:14px; letter-spacing:2px; }
h1,h2,h3 { font-family:Georgia,serif; line-height:1.2; }
h1 { font-size:44px; } h2 { margin-top:50px; border-top:1px solid #c9d2cc; padding-top:25px; }
a { color:#006555; text-underline-offset:3px; } img { max-width:100%; height:auto; }
table { width:100%; border-collapse:collapse; margin:22px 0; font-size:14px; }
th,td { border:1px solid #becdc9; padding:10px; vertical-align:top; overflow-wrap:anywhere; }
th { background:#d8e9e2; text-align:left; } tr { break-inside:avoid; }
pre { background:#172b33; color:#f4f0e7; padding:20px; border-radius:8px; white-space:pre-wrap; overflow-wrap:anywhere; }
code { font-size:.88em; } blockquote { border-left:4px solid #47aa8e; margin:20px 0; padding-left:20px; }
@media(max-width:700px) { main { padding:25px 20px; } h1 { font-size:32px; } table { display:block; overflow:auto; } }
@media print { body { font-size:10pt; } main { padding:0; } h2,h3 { break-after:avoid; } }
"""


def absolute_link(value: str) -> str:
    if urlsplit(value).scheme:
        return value
    base = REPO.replace("/blob/", "/tree/") if (ROOT / urlsplit(value).path).is_dir() else REPO
    return urljoin(base, value)


def inline(paragraph: Paragraph, children: list[Token]) -> None:
    bold = italic = False
    link = ""
    for token in children:
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type == "link_open":
            link = absolute_link(str(token.attrGet("href") or ""))
        elif token.type == "link_close":
            link = ""
        elif token.type in ("softbreak", "hardbreak"):
            paragraph.add_run(" " if token.type == "softbreak" else "\n")
        elif token.type in ("text", "code_inline"):
            if link:
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
                run_element = OxmlElement("w:r")
                properties = OxmlElement("w:rPr")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "007565")
                properties.append(color)
                run_element.append(properties)
                text = OxmlElement("w:t")
                text.text = token.content
                run_element.append(text)
                hyperlink.append(run_element)
                paragraph._p.append(hyperlink)
            else:
                run = paragraph.add_run(token.content)
                run.bold, run.italic = bold, italic
                if token.type == "code_inline":
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)


def build_word(stem: str, tokens: list[Token], output: Path) -> Path:
    doc = Document()
    doc.core_properties.author = "Branchseed team"
    doc.core_properties.title = TITLES[stem]
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.6)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("176A59")
    section.header.paragraphs[0].text = "BRANCHSEED  /  TORALIS LABS HEALTHCARE"
    footer = section.footer.paragraphs[0]
    footer.add_run("Research prototype  •  Fusion release  |  ")
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    doc.add_heading(TITLES[stem], 0)
    doc.add_paragraph("Battle of the Schools • Project copy and team preparation")
    doc.add_picture(str(ROOT / "docs/media/branchseed-cover.png"), width=Cm(17.8))
    doc.add_paragraph("Read the evaluation scope with every metric. Native organizer-Windows timing remains unverified.")
    doc.add_paragraph(f"Source: {REPO}{stem}.md")
    doc.add_page_break()
    index = 0
    lists: list[str] = []
    while index < len(tokens):
        token = tokens[index]
        if token.type == "heading_open":
            level = int(token.tag[1])
            if level > 1:
                inline(doc.add_heading(level=level - 1), tokens[index + 1].children or [])
            index += 2
        elif token.type in ("bullet_list_open", "ordered_list_open"):
            lists.append("List Bullet" if token.type == "bullet_list_open" else "List Number")
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif token.type == "paragraph_open":
            inline(doc.add_paragraph(style=lists[-1] if lists else "Normal"), tokens[index + 1].children or [])
            index += 2
        elif token.type == "fence":
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(token.content.rstrip())
            run.font.name, run.font.size = "Consolas", Pt(8)
        elif token.type == "table_open":
            rows: list[list[list[Token]]] = []
            while tokens[index].type != "table_close":
                part = tokens[index]
                if part.type == "tr_open":
                    rows.append([])
                elif part.type == "inline":
                    rows[-1].append(part.children or [])
                index += 1
            table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for column, children in enumerate(row):
                    paragraph = table.cell(row_index, column).paragraphs[0]
                    inline(paragraph, children)
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
                        if row_index == 0:
                            run.bold = True
            header = table.rows[0]._tr.get_or_add_trPr()
            header.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph()
        index += 1
    destination = output / f"{stem}.docx"
    doc.save(str(destination))
    return destination


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "outputs/fusion-submission-pack")
    parser.add_argument("--pdf", action="store_true", help="Also convert Word documents with LibreOffice.")
    parser.add_argument("--showcase", type=Path, help="Include the separate product showcase MP4.")
    parser.add_argument("--slides", type=Path, help="Include the current live-presentation PDF.")
    parser.add_argument("--deck", type=Path, help="Include the complete deck and render its speaking script.")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    for source, name in ((args.showcase, "branchseed-showcase.mp4"), (args.slides, "branchseed-slides.pdf")):
        if source is not None:
            shutil.copy2(source, args.output / name)
    shutil.copytree(ROOT / "docs/media", args.output / "docs/media", dirs_exist_ok=True)
    shutil.copy2(ROOT / "docs/fusion-restored-validation.json", args.output / "docs/fusion-restored-validation.json")
    sources = {stem: ROOT / f"{stem}.md" for stem in DOCUMENTS}
    if args.deck:
        shutil.copytree(args.deck, args.output / "presentation", dirs_exist_ok=True)
        sources["SPEAKER_SCRIPT"] = args.deck / "SPEAKER_SCRIPT.md"
    shutil.copy2(ROOT / "presentation/SHOWCASE.md", args.output / "SHOWCASE.md")
    shutil.copy2(ROOT / "presentation/SUBMISSION_START.md", args.output / "START_HERE.md")
    markdown = MarkdownIt("commonmark", {"html": True}).enable("table")
    word_paths = []
    for stem, source in sources.items():
        shutil.copy2(source, args.output / source.name)
        tokens = markdown.parse(source.read_text())
        if stem in TITLES:
            word_paths.append(build_word(stem, tokens, args.output))
        seen: dict[str, int] = {}
        for index, token in enumerate(tokens):
            if token.type == "heading_open":
                title = tokens[index + 1].content.replace("`", "")
                slug = re.sub(r"[^\w\- ]", "", title.lower()).replace(" ", "-")
                number = seen.get(slug, 0)
                seen[slug] = number + 1
                token.attrSet("id", f"{slug}-{number}" if number else slug)
            for child in token.children or []:
                if child.type == "link_open":
                    value = str(child.attrGet("href") or "")
                    name, separator, fragment = value.partition("#")
                    if name in {f"{item}.md" for item in sources}:
                        value = name.removesuffix(".md") + ".html" + (separator + fragment if separator else "")
                    elif not value.startswith("#"):
                        value = absolute_link(value)
                    child.attrSet("href", value)
        content = markdown.renderer.render(tokens, markdown.options, {})
        page = (f'<!doctype html><html lang="en"><head><meta charset="utf-8">'
                f'<meta name="viewport" content="width=device-width,initial-scale=1">'
                f'<title>{escape(stem.replace("_", " ").title())} | Branchseed</title>'
                f"<style>{STYLE}</style></head><body><main>"
                f"<header>BRANCHSEED / TORALIS LABS HEALTHCARE</header>{content}</main></body></html>")
        (args.output / f"{stem}.html").write_text(page)
    if args.pdf:
        profile = (ROOT / "outputs/project-documents-office").as_uri()
        subprocess.run(["libreoffice", f"-env:UserInstallation={profile}", "--headless",
                        "--convert-to", "pdf", "--outdir", str(args.output),
                        *map(str, word_paths)], check=True)
        for path in word_paths:
            if not path.with_suffix(".pdf").is_file():
                raise RuntimeError(f"PDF conversion failed: {path.name}")
    links = [
        ("DEVPOST_SUBMISSION.html", "Devpost copy and submission checklist"),
        ("PRESENTER_BRIEFING.html", "Complete presenter briefing and judge Q&A"),
        ("DEMO_GUIDE.html", "Website, Windows and live-demo instructions"),
        ("README.html", "Public project overview and measured results"),
    ]
    for stem, title in TITLES.items():
        if stem not in sources:
            continue
        links.append((f"{stem}.docx", f"{title} — editable Word"))
        if args.pdf:
            links.append((f"{stem}.pdf", f"{title} — PDF"))
    if args.showcase:
        links.append(("branchseed-showcase.mp4", "40-second product showcase — upload separately"))
    if args.slides:
        links.append(("branchseed-slides.pdf", "Five-minute live-demo slides — PDF"))
    if args.deck:
        links.extend([
            ("presentation/branchseed-editable.pptx", "Current fusion presentation — editable PowerPoint"),
            ("presentation/index.html", "Current fusion presentation — offline HTML"),
            ("SPEAKER_SCRIPT.html", "Five-minute script — four speakers"),
        ])
    navigation = "".join(f'<li><a href="{href}">{escape(title)}</a></li>' for href, title in links)
    images = "".join(f'<li><a href="docs/media/{path.name}">{escape(path.stem.replace("-", " ").title())}</a></li>'
                     for path in sorted((args.output / "docs/media").glob("*.png")))
    (args.output / "START_HERE.html").write_text(
        f'<!doctype html><html lang="en"><head><meta charset="utf-8">'
        f'<meta name="viewport" content="width=device-width,initial-scale=1">'
        f"<title>Branchseed submission kit</title><style>{STYLE}</style></head><body><main>"
        f"<header>BRANCHSEED / TORALIS LABS HEALTHCARE</header><h1>Your submission kit</h1>"
        f'<img src="docs/media/branchseed-cover.png" alt="Branchseed — Find the branch. Keep the evidence.">'
        f"<p>Extract the entire ZIP, then open this file. Documents and images work offline; "
        f"source and event links require internet.</p><h2>Prepare and submit</h2><ul>{navigation}</ul>"
        f"<p>The team still needs to submit the Devpost entry, set public video permissions, "
        f"add all teammates, and verify the actual Windows laptop.</p>"
        f"<h2>Cover and gallery</h2><ul>{images}</ul><p>Retain the captions from the submission document. "
        f"Recorded interface views are historical. Reused-reference and synthetic results have different scopes.</p>"
        f'<h2>Technical submission and full editable deck</h2><p><a href="'
        f'https://github.com/Coder-Meet/battleoftheschool/releases/tag/branchseed-judge-fusion-2026-09-13">'
        f"Download the selected fusion judge application and evidence</a>. "
        f"This presentation pack does not include CT volumes or the inference ZIP.</p>"
        f"</main></body></html>"
    )
    checksums = []
    manifest = args.output / "SHA256SUMS"
    for path in sorted(args.output.rglob("*")):
        if path.is_file() and path != manifest:
            with path.open("rb") as stream:
                digest = hashlib.file_digest(stream, "sha256").hexdigest()
            checksums.append(f"{digest}  {path.relative_to(args.output).as_posix()}")
    manifest.write_text("\n".join(checksums) + "\n")
    print(f"Built offline HTML and team documents in {args.output}")


if __name__ == "__main__":
    main()
